# Copyright 2026 Joji John
# SPDX-License-Identifier: Apache-2.0
"""The owner's review: decisions and notes made in LEDGER.xlsx, brought back into charter/ (plan §6d)."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from dev_charter.owner_review import _set_details

from .conftest import FEATURE_F014, Product

F014 = "charter/features/F-014-track-my-order.md"
LIVE = FEATURE_F014.replace("state: agreed", "state: live").replace("decisions: [DEC-001]", "decisions: []")


def three_items(product: Product) -> dict[str, Path]:
    """F-014 agreed and Must; F-015 an idea with no priority; P-001 an idea."""
    product.write(F014, FEATURE_F014.replace("decisions: [DEC-001]", "decisions: []"))
    product.run("feature", "new", "Reset my password")
    product.run("promise", "new", "Customers only see their own orders")
    product.run("report")
    folder = product.root / "charter"
    return {"F-014": product.root / F014, "F-015": next((folder / "features").glob("F-015-*.md")),
            "P-001": next((folder / "promises").glob("P-001-*.md"))}


def edit(path: Path, cells: dict[str, tuple[str, str]]) -> None:
    """Change cells the way a person does in their spreadsheet app: {ID: (decision, note)}, found by header and ID."""
    book = load_workbook(path)
    for sheet in (book["Features"], book["Promises"]):
        header = [c.value for c in sheet[2]]
        decision, note = header.index("My decision") + 1, header.index("My note") + 1
        for row in range(3, sheet.max_row + 1):
            item = sheet.cell(row=row, column=1).value
            if item in cells:
                chosen, written = cells[item]
                if chosen is not None:
                    sheet.cell(row=row, column=decision, value=chosen)
                if written is not None:
                    sheet.cell(row=row, column=note, value=written)
    book.save(path)


def column(path: Path, sheet: str, title: str) -> dict[str, str]:
    ws = load_workbook(path)[sheet]
    at = [c.value for c in ws[2]].index(title)
    return {row[0].value: row[at].value for row in ws.iter_rows(min_row=3)}


def test_the_spreadsheet_opens_with_each_current_decision_and_a_dropdown(product: Product):
    three_items(product)
    ledger = product.root / "LEDGER.xlsx"
    assert column(ledger, "Features", "My decision") == {"F-014": "Agree (Must)", "F-015": "Not reviewed"}
    assert column(ledger, "Features", "My note") == {"F-014": None, "F-015": None}
    assert column(ledger, "Promises", "My decision") == {"P-001": "Not reviewed"}
    book = load_workbook(ledger)
    features, promises = book["Features"], book["Promises"]
    assert features["C3"].fill.fgColor.rgb.endswith("FFF2CC") and features["D4"].fill.fgColor.rgb.endswith("FFF2CC")
    [dropdown] = features.data_validations.dataValidation
    assert dropdown.formula1 == '"Not reviewed,Agree,Agree (Must),Agree (Should),Agree (Could),Defer,Cut"'
    assert str(dropdown.sqref) == "C3:C4"
    [dropdown] = promises.data_validations.dataValidation
    assert dropdown.formula1 == '"Not reviewed,Agree,Defer,Cut"'
    assert "yellow columns are yours" in features["A1"].value
    assert "This tab is generated" in book["Decisions"]["A1"].value


def test_import_applies_only_what_changed_and_leaves_every_other_line_alone(product: Product):
    files = three_items(product)
    before = {item: path.read_text() for item, path in files.items()}
    ledger = product.root / "LEDGER.xlsx"
    edit(ledger, {"F-015": ("Agree (Should)", None), "F-014": (None, "reword check c, merge with F-012"),
                  "P-001": ("Cut", None)})

    output = product.run("import", "--by", "Joji John").output
    assert "F-015 Reset my password: Not reviewed → Agree (Should)" in output
    assert "P-001 Customers only see their own orders: Not reviewed → Cut" in output
    assert "F-014 note: reword check c, merge with F-012" in output

    assert files["F-014"].read_text() == before["F-014"]
    old, new = before["F-015"].splitlines(), files["F-015"].read_text().splitlines()
    changed = [(a, b) for a, b in zip(old, new, strict=True) if a != b]
    assert changed == [('priority: ""', 'priority: "Should"'),
                       ("state: idea                   # idea | agreed | live | parked | retired: changed by people "
                        "only", "state: agreed                 # idea | agreed | live | parked | retired: changed by "
                        "people only")]
    assert "state: retired " in files["P-001"].read_text()

    log = (product.root / "charter/review-log.md").read_text()
    assert "· F-015 · Agree (Should) · Joji John · was Not reviewed" in log
    assert "· P-001 · Cut · Joji John · was Not reviewed" in log
    notes = (product.root / "charter/review-notes.md").read_text()
    assert notes.rstrip().endswith("· F-014 · Joji John · reword check c, merge with F-012")

    assert column(ledger, "Features", "My decision") == {"F-014": "Agree (Must)", "F-015": "Agree (Should)"}
    assert column(ledger, "Features", "My note") == {"F-014": None, "F-015": None}
    assert column(ledger, "Promises", "Status") == {"P-001": "Retired"}
    assert "Nothing to import" in product.run("import").output


def test_a_preview_changes_nothing(product: Product):
    three_items(product)
    edit(product.root / "LEDGER.xlsx", {"F-015": ("Defer", "later")})
    before = {p: p.read_bytes() for p in (product.root / "charter").rglob("*") if p.is_file()}
    output = product.run("import", "--preview").output
    assert "Not reviewed → Defer" in output and "Preview only" in output
    assert {p: p.read_bytes() for p in (product.root / "charter").rglob("*") if p.is_file()} == before


def test_nothing_is_imported_while_any_cell_cannot_be_applied(product: Product):
    files = three_items(product)
    files["F-014"].write_text(LIVE)
    product.run("report")
    edit(product.root / "LEDGER.xlsx", {"F-014": ("Defer", None), "F-015": ("Maybe", None), "P-001": ("Cut", None)})
    before = {item: path.read_text() for item, path in files.items()}
    result = product.run("import", "--by", "Joji John", ok=False)
    assert result.exit_code == 1 and "Nothing was imported" in result.output
    assert "F-014 (Features tab, row 3) is live, so it cannot be deferred" in result.output
    assert "'Maybe' is not one of the choices" in result.output
    assert {item: path.read_text() for item, path in files.items()} == before
    assert not (product.root / "charter/review-log.md").exists()


def test_agree_keeps_a_live_feature_live_and_sets_its_priority(product: Product):
    files = three_items(product)
    files["F-014"].write_text(LIVE)
    product.run("report")
    edit(product.root / "LEDGER.xlsx", {"F-014": ("Agree (Could)", None)})
    assert "Agree (Must) → Agree (Could)" in product.run("import", "--by", "Joji John").output
    text = files["F-014"].read_text()
    assert "state: live" in text and "priority: Could" in text


def test_a_copy_imports_however_it_was_rearranged_and_the_local_review_is_kept(product: Product, tmp_path: Path):
    three_items(product)
    ledger = product.root / "LEDGER.xlsx"
    edit(ledger, {"F-014": (None, "a note of my own")})
    kept = ledger.read_bytes()

    copy = Workbook()
    sheet = copy.active
    sheet.title = "From Google Sheets"
    sheet.append(["Title", "My Decision", "ID", "my note"])  # no note row; columns moved; header case changed
    sheet.append(["Reset my password", "agree – must", "F-015", ""])
    sheet.append(["Customers only see their own orders", "  CUT ", "P-001", ""])
    sheet.append(["Gone", "Agree", "F-099", ""])
    path = tmp_path / "copy.xlsx"
    copy.save(path)

    output = product.run("import", str(path), "--by", "Joji John").output
    assert "Skipped F-099" in output
    assert "Not reviewed → Agree (Must)" in output and "Not reviewed → Cut" in output
    assert "was not rebuilt, so your review in it is kept: it holds 1 note" in output
    assert ledger.read_bytes() == kept


def test_a_note_on_an_item_no_longer_in_the_charter_is_refused(product: Product, tmp_path: Path):
    three_items(product)
    copy = Workbook()
    copy.active.append(["ID", "My decision", "My note"])
    copy.active.append(["F-099", "", "keep this"])
    path = tmp_path / "copy.xlsx"
    copy.save(path)
    result = product.run("import", str(path), ok=False)
    assert "F-099" in result.output and "no longer in charter/" in result.output


def test_two_rows_for_one_item_must_agree(product: Product, tmp_path: Path):
    three_items(product)
    copy = Workbook()
    for row in (["ID", "My decision"], ["F-015", "Defer"], ["F-015", "Cut"]):
        copy.active.append(row)
    path = tmp_path / "copy.xlsx"
    copy.save(path)
    assert "appears twice with different decisions" in product.run("import", str(path), ok=False).output


def test_a_note_is_not_brought_in_twice(product: Product, tmp_path: Path):
    three_items(product)
    copy = Workbook()
    for row in (["ID", "My decision", "My note"], ["F-015", "", "needs a design first"]):
        copy.active.append(row)
    path = tmp_path / "copy.xlsx"
    copy.save(path)
    product.run("import", str(path), "--by", "Ana")
    product.run("import", str(path), "--by", "Ana")
    notes = (product.root / "charter/review-notes.md").read_text()
    assert notes.count("needs a design first") == 1
    assert "1 review note waiting in charter/review-notes.md" in product.run("status", "--brief").output


def test_the_rebuild_after_a_commit_keeps_an_unimported_review(git_product: Product):
    git_product.write(F014, FEATURE_F014.replace("decisions: [DEC-001]", "decisions: []"))
    git_product.commit("feature")
    ledger, page = git_product.root / "LEDGER.xlsx", git_product.root / "status.html"
    assert ledger.is_file(), "the post-commit hook builds the spreadsheet"
    edit(ledger, {"F-014": ("Agree (Should)", "check c needs a real phone")})
    kept, old_page = ledger.read_bytes(), page.read_text()

    git_product.run("park", "Show the courier on a map")
    git_product.commit("an idea")
    assert ledger.read_bytes() == kept
    assert page.read_text() != old_page, "status.html is still rebuilt"
    assert "LEDGER.xlsx was not rebuilt" in git_product.run("report").output
    assert "holds a review not imported yet (1 decision and 1 note)" in git_product.run("status", "--brief").output

    git_product.run("import", "--by", "Joji John")
    assert column(ledger, "Features", "My note") == {"F-014": None}
    assert "Rebuilt LEDGER.xlsx" in git_product.run("report").output


def test_an_unreadable_spreadsheet_is_kept_and_explained(product: Product):
    three_items(product)
    ledger = product.root / "LEDGER.xlsx"
    ledger.write_bytes(b"")
    result = product.run("report", ok=False)
    assert "LEDGER.xlsx was not rebuilt" in result.output and "delete it" in result.output
    assert ledger.read_bytes() == b""


def test_details_keep_their_quoting_comment_column_and_line_endings():
    text = "---\r\nid: F-001\r\npriority: \"\"\r\nstate: idea      # people only\r\n---\r\nbody\r\n"
    assert _set_details(text, {"state": "agreed", "priority": "Must"}) == (
        "---\r\nid: F-001\r\npriority: \"Must\"\r\nstate: agreed    # people only\r\n---\r\nbody\r\n")
    assert _set_details("---\nid: F-001\nstate: idea\n---\n", {"priority": "Could"}) == (
        "---\nid: F-001\nstate: idea\npriority: Could\n---\n")


def test_a_file_that_records_the_wrong_size_still_gives_every_row(tmp_path: Path):
    import re
    import zipfile

    from dev_charter.owner_review import read_sheet

    book = Workbook()
    for row in (["ID", "My decision", "My note"], ["F-001", "Cut", ""], ["F-002", "Defer", "later"]):
        book.active.append(row)
    book.save(tmp_path / "ok.xlsx")
    with zipfile.ZipFile(tmp_path / "ok.xlsx") as source, zipfile.ZipFile(tmp_path / "bad.xlsx", "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data)
            target.writestr(item, data)
    assert [row.id for row in read_sheet(tmp_path / "bad.xlsx")] == ["F-001", "F-002"]
