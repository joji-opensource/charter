# Copyright 2026 Joji John
# SPDX-License-Identifier: Apache-2.0
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from dev_charter import git
from dev_charter.config import Repo
from dev_charter.results import ids_in
from dev_charter.scan import is_code, is_test, read_tags

from .conftest import FEATURE_F014, Product

DECISION = FEATURE_F014.replace("decisions: [DEC-001]", "decisions: []")


@pytest.mark.parametrize(
    ("text", "refs", "ids", "foundation", "line", "words"),
    [
        ('"""Order status."""\n# serves: F-014\n', ("F-014",), ("F-014",), False, 2, ""),
        ("// serves: F-014, P-002\n", ("F-014", "P-002"), ("F-014", "P-002"), False, 1, ""),
        ("# serves: F-014.a\n", ("F-014.a",), ("F-014",), False, 1, ""),
        ("# serves: foundation\n", (), (), True, 1, ""),
        ("# serves: F-014 — Track my order\n", ("F-014",), ("F-014",), False, 1, "Track my order"),
        ("/* serves: F-014.c - someone else's order asks to sign in */\n", ("F-014.c",), ("F-014",), False, 1,
         "someone else's order asks to sign in"),
    ],
)
def test_a_tag_names_what_the_code_serves(text, refs, ids, foundation, line, words):
    (tag,) = read_tags(text)
    assert (tag.refs, tag.ids, tag.foundation, tag.line, tag.words) == (refs, ids, foundation, line, words)


def test_every_tag_in_a_file_counts():
    text = ("# serves: F-014 — Track my order\n\n\ndef status():\n    pass\n\n\n"
            "# serves: F-017.a — gift wrap at checkout\ndef wrap():\n    pass\n")
    tags = read_tags(text)
    assert [(t.refs, t.line) for t in tags] == [(("F-014",), 1), (("F-017.a",), 8)]


@pytest.mark.parametrize(
    "text",
    [
        "the step this serves: the person's own words where the document traces it",  # prose found in P4
        "from x import y  # noqa: F401\ncolor = '#f800'\n",  # linter codes and colours found in P4
        "# serves: F-14\n",
    ],
)
def test_prose_and_look_alikes_are_not_tags(text):
    assert read_tags(text) == ()


@pytest.mark.parametrize(
    ("path", "code", "test"),
    [
        ("shop/orders/status.py", True, False),
        ("tests/test_status.py", False, True),
        ("web/src/order.test.ts", False, True),
        ("pkg/status_test.go", False, True),
        ("node_modules/lib/index.js", False, False),
        ("docs/guide.md", False, False),
        ("charter/features/F-014-x.md", False, False),
    ],
)
def test_code_and_tests_are_told_apart(path, code, test):
    repo = Repo(path=Path("."))
    assert (is_code(path, repo), is_test(path, repo)) == (code, test)


def test_patterns_in_charter_toml_override_the_defaults():
    repo = Repo(path=Path("."), code=("app/**/*.py",), tests=("checks/**",))
    assert is_code("app/orders/status.py", repo) and not is_code("scripts/deploy.py", repo)
    assert is_test("checks/orders.py", repo) and not is_code("checks/orders.py", repo)


@pytest.mark.parametrize(
    ("name", "ids"),
    [
        ("test_F014_b_text_sent_when_out_for_delivery", ("F-014.b",)),
        ("F-014.b sends a text when the order goes out", ("F-014.b",)),
        ("TestF014B", ("F-014.b",)),
        ("test_P002_a_charged_amount_matches", ("P-002.a",)),
        ("test_F014_order_page", ("F-014",)),
        ("test_vehicle_names[F-150-truck]", ()),  # parameters are not IDs (P3)
        ("test_f401_lint_rule", ()),
        ("testF014bSendsText", ("F-014.b",)),
        ("test_F014_order_page_loads", ("F-014",)),
    ],
)
def test_test_names_carry_their_check(name, ids):
    assert ids_in(name) == ids


@pytest.mark.parametrize(
    ("remote", "url"),
    [
        ("git@github.com:joji-opensource/charter.git", "https://github.com/joji-opensource/charter"),
        ("https://github.com/joji-opensource/charter.git", "https://github.com/joji-opensource/charter"),
        ("https://token@gitlab.com/team/shop", "https://gitlab.com/team/shop"),
        ("ssh://git@bitbucket.org/team/shop.git", "https://bitbucket.org/team/shop"),
    ],
)
def test_the_web_address_comes_from_the_remote(tmp_path, remote, url):
    subprocess.run(["git", "init", "-q", str(tmp_path)], check=True)
    subprocess.run(["git", "-C", str(tmp_path), "remote", "add", "origin", remote], check=True)
    assert git.web_url(tmp_path) == url


def test_links_follow_each_host():
    assert git.file_link("https://github.com/o/r", "abc", "a/b.py") == "https://github.com/o/r/blob/abc/a/b.py"
    assert git.file_link("https://gitlab.com/o/r", "abc", "a/b.py") == "https://gitlab.com/o/r/-/blob/abc/a/b.py"


JUNIT = """<?xml version="1.0" encoding="utf-8"?>
<testsuites><testsuite name="pytest" tests="3">
  <testcase classname="tests.test_orders" name="test_F014_a_shows_step_and_date"/>
  <testcase classname="tests.test_orders" name="test_F014_c_other_customers_order_signs_in">{c}</testcase>
  <testcase classname="tests.test_orders" name="test_F099_a_not_in_the_charter"/>
</testsuite></testsuites>
"""


def test_one_red_test_turns_a_proven_feature_at_risk(git_product: Product):
    shop = git_product
    shop.write("charter/features/F-014-track-my-order.md", DECISION + (
        "- 2026-10-24 · proven · Joji John · watched the whole flow work on staging\n"))
    shop.write("shop/orders/status.py", "# serves: F-014\ndef step(order):\n    return order['step']\n")
    shop.write("shop/money.py", "# serves: foundation\n")
    shop.write("shop/legacy.py", "def old():\n    pass\n")
    shop.set_results("reports/*.xml")
    shop.write("reports/junit.xml", JUNIT.format(c=""))
    shop.commit("feature, and code from before Dev Charter", check=False)
    subprocess.run(["git", "-C", str(shop.root), "update-ref", "refs/remotes/origin/main", "HEAD"], check=True)

    shop.run("report")
    book = load_workbook(shop.root / "LEDGER.xlsx")
    features = book["Features"]
    assert (features["F3"].value, features["G3"].value) == ("Proven", "3 of 3 checks passing")
    assert features["H3"].value == "shop: shop/orders/status.py"
    assert features["I3"].value.startswith("2 automated checks (2 passed)")
    code_map = {row[1].value: row for row in book["Code map"].iter_rows(min_row=3)}
    assert [code_map[p][2].value for p in ("shop/legacy.py", "shop/money.py", "shop/orders/status.py")] == [
        "Unexplained", "Shared foundation", "Feature code"]
    head = git.head(shop.root)
    assert code_map["shop/orders/status.py"][1].hyperlink.target == (
        f"https://github.com/acme/shop/blob/{head}/shop/orders/status.py")

    shop.write("reports/junit.xml", JUNIT.format(c='<failure message="opened another order"/>'))
    shop.run("report")
    features = load_workbook(shop.root / "LEDGER.xlsx")["Features"]
    assert (features["F3"].value, features["G3"].value) == ("At risk", "2 of 3 checks passing")
    assert "✗ c." in features["O3"].value

    output = shop.run("check").output
    assert "test_F099_a_not_in_the_charter' names a feature or promise that is not in charter/" in output
    assert "1 other code file(s) do not say which feature they serve" in output


def test_an_unknown_feature_in_a_tag_is_a_problem(git_product: Product):
    git_product.write("shop/orders/status.py", "# serves: F-099\n")
    result = git_product.run("check", ok=False)
    assert result.exit_code == 1
    assert "serves F-099, but F-099 is not a feature or promise" in result.output
    assert "shop/orders/status.py:1" in result.output
