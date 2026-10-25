# Copyright 2026 Joji John
# SPDX-License-Identifier: Apache-2.0
"""The owner's review: the My decision and My note columns of LEDGER.xlsx, brought back into charter/.

LEDGER.xlsx is generated, but those two columns belong to the person who owns the product. Each decision opens
showing what the files say now; `dev-charter import` applies the ones they changed and collects their notes, and the
report never replaces a spreadsheet holding a review that has not been imported.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from .config import Config
from .errors import CharterError, FileProblem
from .model import FEATURE, PRIORITIES, Entry, Kind
from .store import Project, read_entry

DECISION_HEADER, NOTE_HEADER = "My decision", "My note"
NOT_REVIEWED, AGREE, DEFER, CUT = "Not reviewed", "Agree", "Defer", "Cut"
NOTES_FILE, LOG_FILE = "review-notes.md", "review-log.md"
ID_RE = re.compile(r"^[FP]-\d{3,}$")
NOTES_HEADER = """# Review notes

Notes written in the My note column of LEDGER.xlsx, brought in by `dev-charter import`. Go through each one with
the person who wrote it, change nothing it asks for until they agree, and delete it once it is dealt with.

"""
LOG_HEADER = """# Review log

Decisions brought in from the My decision column of LEDGER.xlsx by `dev-charter import`:
date · item · decision · who decided · what it was before. Never edit or delete a line.

"""


def choices(kind: Kind) -> tuple[str, ...]:
    """What the dropdown offers: promises have no priority."""
    if kind is FEATURE:
        return (NOT_REVIEWED, AGREE, *(f"{AGREE} ({p})" for p in PRIORITIES), DEFER, CUT)
    return (NOT_REVIEWED, AGREE, DEFER, CUT)


def current(entry: Entry) -> str:
    """The decision the files hold now, as the spreadsheet shows it."""
    if entry.state in ("agreed", "live"):
        return f"{AGREE} ({entry.priority})" if entry.priority and entry.kind is FEATURE else AGREE
    return {"parked": DEFER, "retired": CUT}.get(entry.state, NOT_REVIEWED)


def _key(text: str) -> str:
    return re.sub(r"[^a-z]", "", text.lower())


def _choice(text: str, kind: Kind) -> str | None:
    """The choice a cell names, ignoring case, spaces, brackets and dashes; None if it names none."""
    return next((c for c in choices(kind) if _key(c) == _key(text)), None)


# --- reading a spreadsheet -------------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Row:
    sheet: str
    line: int
    id: str
    decision: str
    note: str

    @property
    def where(self) -> str:
        return f"{self.sheet} tab, row {self.line}"


def _text(values: tuple, index: int | None) -> str:
    if index is None or index >= len(values) or values[index] is None:
        return ""
    return " ".join(str(values[index]).split())


def read_sheet(path: Path) -> list[Row]:
    """Every row with an ID on any tab whose header has ID and My decision. Columns are found by their header and rows
    by their ID, so sorting, filtering or moving columns in a spreadsheet app changes nothing."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many kinds of error for a file it cannot read
        raise CharterError(f"{path.name} could not be read as a spreadsheet ({exc}); it must be an Excel workbook "
                           "(.xlsx).") from exc
    rows: list[Row] = []
    try:
        for sheet in workbook.worksheets:
            sheet.reset_dimensions()  # some apps record a wrong size, which would cut rows off
            header = None
            for number, values in enumerate(sheet.iter_rows(max_row=10, values_only=True), start=1):
                names = [_text(values, i).casefold() for i in range(len(values))]
                if "id" in names and DECISION_HEADER.casefold() in names:
                    header = number, names
                    break
            if header is None:
                continue
            first, names = header
            at_id, at_decision = names.index("id"), names.index(DECISION_HEADER.casefold())
            at_note = names.index(NOTE_HEADER.casefold()) if NOTE_HEADER.casefold() in names else None
            for number, values in enumerate(sheet.iter_rows(min_row=first + 1, values_only=True), start=first + 1):
                item = _text(values, at_id)
                if ID_RE.match(item):
                    rows.append(Row(sheet.title, number, item, _text(values, at_decision), _text(values, at_note)))
    finally:
        workbook.close()
    return rows


# --- working out what changes ----------------------------------------------------------------------------------------


@dataclass(frozen=True)
class Change:
    entry: Entry
    choice: str
    state: str
    priority: str

    @property
    def before(self) -> str:
        return current(self.entry)


@dataclass
class Review:
    """What a spreadsheet asks for, compared with the files."""

    changes: list[Change] = field(default_factory=list)
    notes: list[tuple[Entry, str]] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)  # cells that cannot be applied; nothing is, while any remain
    skipped: list[str] = field(default_factory=list)  # IDs no longer in charter/, with no note

    @property
    def waiting(self) -> bool:
        """Whether the spreadsheet holds anything not yet imported."""
        return bool(self.changes or self.notes or self.problems)

    def summary(self) -> str:
        parts = []
        if self.changes:
            parts.append(f"{len(self.changes)} decision{'s' if len(self.changes) != 1 else ''}")
        if self.notes:
            parts.append(f"{len(self.notes)} note{'s' if len(self.notes) != 1 else ''}")
        if self.problems:
            parts.append(f"{len(self.problems)} cell{'s' if len(self.problems) != 1 else ''} to correct")
        return " and ".join(parts)


def _outcome(entry: Entry, choice: str) -> tuple[str, str]:
    """The state and priority a choice asks for."""
    if choice == CUT:
        state = "retired"
    elif choice == DEFER:
        state = "parked"
    else:
        state = "live" if entry.state == "live" else "agreed"
    priority = choice[len(AGREE) + 2:-1] if choice.startswith(f"{AGREE} (") else entry.priority
    return state, priority


def plan(project: Project, rows: list[Row]) -> Review:
    entries = {e.id: e for e in project.entries}
    review = Review()
    chosen: dict[str, tuple[str, Row]] = {}
    for row in rows:
        entry = entries.get(row.id)
        if entry is None:
            if row.note:
                review.problems.append(f"{row.id} ({row.where}) has a note, but {row.id} is no longer in charter/. "
                                       "Move the note to the item it belongs to, or clear it.")
            elif row.id not in review.skipped:
                review.skipped.append(row.id)
            continue
        if row.note and (entry, row.note) not in review.notes:
            review.notes.append((entry, row.note))
        if not row.decision:
            continue
        choice = _choice(row.decision, entry.kind)
        if choice is None:
            review.problems.append(f"{row.id} ({row.where}): '{row.decision}' is not one of the choices: "
                                   f"{', '.join(choices(entry.kind))}.")
        elif choice != NOT_REVIEWED:
            earlier = chosen.get(entry.id)
            if earlier and earlier[0] != choice:
                review.problems.append(f"{row.id} appears twice with different decisions: {earlier[0]} "
                                       f"({earlier[1].where}) and {choice} ({row.where}). Keep one.")
            chosen[entry.id] = choice, row
    for entry_id, (choice, row) in chosen.items():
        entry = entries[entry_id]
        if choice == DEFER and entry.state == "live":
            review.problems.append(f"{entry_id} ({row.where}) is live, so it cannot be deferred. Choose Cut to "
                                   f"withdraw it, or leave it as {current(entry)}.")
            continue
        state, priority = _outcome(entry, choice)
        if (state, priority) != (entry.state, entry.priority):
            review.changes.append(Change(entry, choice, state, priority))
    return review


def unimported(project: Project, path: Path) -> Review | None:
    """The review a spreadsheet holds that has not been imported yet; None when there is none, or no spreadsheet."""
    if not path.is_file():
        return None
    review = plan(project, read_sheet(path))
    return review if review.waiting else None


# --- applying it -----------------------------------------------------------------------------------------------------


def _set_details(text: str, values: dict[str, str]) -> str:
    """Set lines in a file's details block, keeping each line's quoting and comment (and the comment's column)."""
    lines = text.splitlines(keepends=True)
    end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
    for key, value in values.items():
        pattern = re.compile(rf"^{key}:(?P<gap>[ \t]*)(?P<value>.*?)(?P<pad>[ \t]*)(?P<comment>#.*)?(?P<eol>\r?\n?)$")
        index = next((i for i in range(1, end) if pattern.match(lines[i])), None)
        if index is None:
            eol = "\r\n" if lines[end].endswith("\r\n") else "\n"
            lines.insert(end, f"{key}: {value}{eol}")
            end += 1
            continue
        match = pattern.match(lines[index])
        old = match.group("value")
        new = json.dumps(value) if old[:1] in ('"', "'") or not value else value
        pad = match.group("pad")
        if match.group("comment"):
            pad = " " * max(1, len(old) + len(pad) - len(new))
        lines[index] = (f"{key}:{match.group('gap') or ' '}{new}{pad}{match.group('comment') or ''}"
                        f"{match.group('eol')}")
    return "".join(lines)


def _append(path: Path, header: str, lines: list[str]) -> None:
    text = path.read_text(encoding="utf-8") if path.exists() else header
    if text and not text.endswith("\n"):
        text += "\n"
    path.write_text(text + "".join(f"{line}\n" for line in lines), encoding="utf-8")


def open_notes(config: Config) -> int:
    path = config.charter_dir / NOTES_FILE
    if not path.is_file():
        return 0
    return sum(1 for line in path.read_text(encoding="utf-8").splitlines() if line.startswith("- "))


def apply(config: Config, review: Review, person: str, today: date | None = None) -> list[Path]:
    """Write the decisions into the files, log them, and add the notes. Every new file is checked before any is
    written, so either all of it applies or none of it does. Returns the files changed."""
    if review.problems:
        raise CharterError("Nothing was imported: some cells cannot be applied.")
    today = today or date.today()
    person = " ".join(person.replace("·", "-").split())
    rewritten: list[tuple[Path, str]] = []
    for change in review.changes:
        path = change.entry.path
        values = {"state": change.state}
        if change.priority != change.entry.priority:
            values["priority"] = change.priority
        text = _set_details(path.read_text(encoding="utf-8"), values)
        try:
            entry = read_entry(path, change.entry.kind, text)
        except FileProblem as problem:
            raise CharterError(f"Nothing was imported: {path.name} would no longer read correctly "
                               f"({problem.message}).") from problem
        if (entry.state, entry.priority) != (change.state, change.priority):
            raise CharterError(f"Nothing was imported: {path.name} could not be changed safely; change it by hand.")
        rewritten.append((path, text))
    for path, text in rewritten:
        path.write_text(text, encoding="utf-8")
    changed = [path for path, _ in rewritten]

    if review.changes:
        log = config.charter_dir / LOG_FILE
        _append(log, LOG_HEADER, [f"- {today.isoformat()} · {c.entry.id} · {c.choice} · {person} · was {c.before}"
                                  for c in review.changes])
        changed.append(log)
    notes_path = config.charter_dir / NOTES_FILE
    existing = notes_path.read_text(encoding="utf-8").splitlines() if notes_path.exists() else []
    new_notes = []
    for entry, note in review.notes:
        note = note.replace("·", "-")
        if not any(f" · {entry.id} · " in line and line.endswith(f" · {note}") for line in existing):
            new_notes.append(f"- {today.isoformat()} · {entry.id} · {person} · {note}")
    if new_notes:
        _append(notes_path, NOTES_HEADER, new_notes)
        changed.append(notes_path)
    return changed
