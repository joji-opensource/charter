# Copyright 2026 Joji John
# SPDX-License-Identifier: Apache-2.0
from pathlib import Path

from click.testing import CliRunner
from openpyxl import load_workbook

from dev_charter import __version__, guides
from dev_charter.cli import main

from .conftest import DECISION_001, FEATURE_F014, Product


def snapshot(root: Path) -> dict[str, bytes]:
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in sorted(root.rglob("*")) if p.is_file()}


def test_init_sets_everything_up_and_a_second_run_changes_nothing(product):
    root = product.root
    for rel in ("charter.toml", "AGENTS.md", ".gitignore", ".gitattributes", "charter/ideas.md",
                "charter/features/.gitkeep", "charter/cards/.gitkeep"):
        assert (root / rel).is_file(), rel
    assert not (root / "charter/handoff.md").exists()
    assert "/LEDGER.xlsx" in (root / ".gitignore").read_text()
    assert "/.dev-charter/" in (root / ".gitignore").read_text()
    assert "charter/ideas.md merge=union" in (root / ".gitattributes").read_text()
    before = snapshot(root)
    result = product.run("init", "--yes")
    assert "already there" in result.output
    assert snapshot(root) == before


def test_init_refuses_a_folder_inside_another_product(product):
    nested = product.root / "app"
    nested.mkdir()
    result = CliRunner().invoke(main, ["--root", str(nested), "init", "--yes"])
    assert result.exit_code == 1 and "inside another Dev Charter product" in result.output


def test_init_refuses_a_repository_that_does_not_exist(tmp_path):
    result = CliRunner().invoke(main, ["--root", str(tmp_path), "init", "--repo", "missing", "--yes"])
    assert result.exit_code == 1 and "'missing' does not exist" in result.output


def test_commands_outside_a_product_say_how_to_start(tmp_path):
    result = CliRunner().invoke(main, ["--root", str(tmp_path), "status"])
    assert result.exit_code == 1 and "Run `dev-charter init`" in result.output


def test_a_project_set_up_for_another_version_is_refused(product):
    toml = product.root / "charter.toml"
    toml.write_text(toml.read_text().replace('dev-charter = ">=0.1,<0.2"', 'dev-charter = ">=2.0,<3"'))
    result = product.run("status", ok=False)
    assert result.exit_code == 1 and "uses Dev Charter >=2.0,<3" in result.output


def test_new_items_get_the_next_ids(product):
    assert "Created F-001" in product.run("feature", "new", "Track my order").output
    assert "Created F-002" in product.run("feature", "new", "Reset my password", "--priority", "Should").output
    assert "Created P-001" in product.run("promise", "new", "Customers only see their own orders").output
    assert "Created DEC-001" in product.run("decision", "new", "Dates from the carrier", "--feature", "F-001").output
    assert "Created CARD-001" in product.run("card", "new", "Order page", "--serves", "F-001.a").output
    bad = product.run("card", "new", "Order page", "--serves", "order page", ok=False)
    assert bad.exit_code == 1 and "is not a feature, promise or check ID" in bad.output


def test_check_reports_problems_with_file_and_line_and_fails(product):
    product.write("charter/features/F-014-track-my-order.md", FEATURE_F014)
    result = product.run("check", ok=False)
    assert result.exit_code == 1
    assert "F-014 refers to decision DEC-001, which does not exist." in result.output
    product.write("charter/decisions/DEC-001-dates-from-the-carrier.md", DECISION_001)
    result = product.run("check")
    assert "0 problem(s) to fix" in result.output


def test_check_warns_about_in_house_words(product):
    toml = product.root / "charter.toml"
    toml.write_text(toml.read_text().replace("internal = []", 'internal = ["SKU"]'))
    product.write("charter/features/F-014-track-my-order.md",
                  FEATURE_F014.replace("see where my order is", "see where each SKU is"))
    product.write("charter/decisions/DEC-001-dates-from-the-carrier.md", DECISION_001)
    assert "uses the in-house word 'SKU'" in product.run("check").output


def test_report_builds_a_readable_spreadsheet_and_page(product):
    product.write("charter/features/F-014-track-my-order.md", FEATURE_F014)
    product.write("charter/decisions/DEC-001-dates-from-the-carrier.md", DECISION_001)
    product.run("park", "Show the courier on a map", "--from", "F-014")
    assert "Rebuilt LEDGER.xlsx and status.html." in product.run("report").output

    book = load_workbook(product.root / "LEDGER.xlsx")
    assert book.sheetnames == ["Summary", "Features", "Promises", "Code map", "Work in progress", "Decisions",
                               "Ideas"]
    features = book["Features"]
    assert features["A1"].value.startswith("Generated ") and "Do not edit" in features["A1"].value
    assert [c.value for c in features[2]][:7] == ["ID", "Title", "Area", "Priority", "Release", "Status", "Checks"]
    assert features["A3"].value == "F-014" and features["F3"].value == "Agreed"
    assert features["G3"].value == "1 of 3 checks passing"
    assert [c.value for c in features[2]][7:9] == ["Built in", "Proven by"]
    assert features["I3"].value == "1 sign-off by Joji John"
    assert features["O3"].value.splitlines()[1].startswith("✓ b.")
    assert features.freeze_panes == "B3"
    assert features.row_dimensions[3].height > 15 * 3
    assert features["F3"].fill.fgColor.rgb.endswith("E7E6E6")
    assert book["Ideas"]["B3"].value == "Show the courier on a map"

    page = (product.root / "status.html").read_text()
    assert "Track my order" in page and "1 of 3 checks passing" in page


def test_report_refuses_rather_than_leave_a_feature_out(product):
    product.write("charter/features/F-014-track-my-order.md", FEATURE_F014.replace("state: agreed", "state: ?"))
    result = product.run("report", ok=False)
    assert result.exit_code == 1
    assert "charter/features/F-014-track-my-order.md:7" in result.output
    assert not (product.root / "LEDGER.xlsx").exists()


def test_status_brief_gives_the_counts_and_each_cards_next_step(product):
    product.write("charter/features/F-014-track-my-order.md", FEATURE_F014)
    product.run("card", "new", "Order page", "--serves", "F-014.a")
    card = next((product.root / "charter/cards").glob("CARD-001-*.md"))
    card.write_text(card.read_text().replace("- Next step:", "- Next step: add the delivery date to the page"))
    output = product.run("status", "--brief").output
    assert "Demo Store · mode standard" in output
    assert "Features: 1 (1 Agreed)" in output
    assert "Open card CARD-001 Order page (F-014.a): next step: add the delivery date to the page" in output


def test_a_leftover_handoff_file_is_pointed_at(product):
    product.write("charter/handoff.md", "# Handoff\n")
    assert "handoff.md is no longer used" in product.run("check").output


def test_every_guide_is_packaged_and_listed():
    runner = CliRunner()
    listing = runner.invoke(main, ["guide"]).output
    for stage in guides.STAGES:
        assert stage in listing
        text = runner.invoke(main, ["guide", stage]).output
        assert text.startswith("# ") and len(text) > 500
    unknown = runner.invoke(main, ["guide", "deploy"])
    assert unknown.exit_code == 1 and "no guide called 'deploy'" in unknown.output


def test_check_notices_rules_that_changed_within_a_minor_release(product):
    agents_md = product.root / "AGENTS.md"
    agents_md.write_text(agents_md.read_text().replace(f"v{__version__} -->", "v0.1 -->", 1))
    assert "its rules have changed since" in product.run("check").output
    product.run("sync")
    assert "Dev Charter section" not in product.run("check").output


def test_sync_refreshes_the_agents_block(product):
    agents_md = product.root / "AGENTS.md"
    agents_md.write_text(agents_md.read_text().replace("Never write an ID yourself.", "old wording"))
    assert "AGENTS.md: updated." in product.run("sync").output
    assert "Never write an ID yourself." in agents_md.read_text()


def test_product_fixture_root(product: Product):
    assert (product.root / "charter.toml").read_text().count("[[repos]]") == 1
