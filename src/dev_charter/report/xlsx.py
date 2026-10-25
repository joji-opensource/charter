# Copyright 2026 Joji John
# SPDX-License-Identifier: Apache-2.0
"""LEDGER.xlsx: the charter as a spreadsheet a product owner can read in five minutes."""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..model import FEATURE, PROMISE, Entry
from ..owner_review import DECISION_HEADER, NOTE_HEADER, choices, current
from ..status import FAILING, ORDER, PASSING
from . import Snapshot

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
YOURS_HEADER_FILL = PatternFill("solid", fgColor="BF8F00")
YOURS_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE_FONT = Font(italic=True, color="595959")
TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=12)
LINK_FONT = Font(color="0563C1", underline="single")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_STYLE = {
    "At risk": ("F8CBAD", "000000"),
    "Live": ("548235", "FFFFFF"),
    "Proven": ("A9D08E", "000000"),
    "Built": ("BDD7EE", "000000"),
    "Being built": ("FFE699", "000000"),
    "Agreed": ("E7E6E6", "000000"),
    "Idea": ("FFFFFF", "000000"),
    "Parked": ("F2F2F2", "7F7F7F"),
    "Retired": ("F2F2F2", "7F7F7F"),
}
CODE_STYLE = {
    "Feature code": ("FFFFFF", "000000"),
    "Shared foundation": ("E7E6E6", "000000"),
    "Unexplained": ("F8CBAD", "000000"),
}
CHECK_MARK = {PASSING: "✓", FAILING: "✗"}
LINE_HEIGHT = 15
MAX_ROW_HEIGHT = 409


@dataclass
class Column:
    title: str
    width: int


def _do_not_edit(snap: Snapshot) -> str:
    return (f"{snap.provenance} This tab is generated: change the files in charter/ (or ask your coding agent) "
            "instead. Your decisions and notes go in the yellow columns of the Features and Promises tabs.")


def _yours(snap: Snapshot) -> str:
    return (f"{snap.provenance} The yellow columns are yours: choose a decision, write a note, save, then ask your "
            "coding agent to import your review (or run dev-charter import). Everything else comes from the files in "
            "charter/: change those, or ask your coding agent.")


def _editable(ws: Worksheet, decision_col: int, note_col: int, options: tuple[str, ...], count: int) -> None:
    """Colour the owner's two columns and give the decision a dropdown of `options`."""
    for col in (decision_col, note_col):
        ws.cell(row=2, column=col).fill = YOURS_HEADER_FILL
        for row in range(3, count + 3):
            ws.cell(row=row, column=col).fill = YOURS_FILL
    if not count:
        return
    listed = ", ".join(options)
    validation = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Choose from the list",
                                error=f"Choose one of: {listed}.")
    letter = get_column_letter(decision_col)
    validation.add(f"{letter}3:{letter}{count + 2}")
    ws.add_data_validation(validation)


def _styled(cell, style: tuple[str, str]) -> None:
    fill, font = style
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(bold=True, color=font)


def _status_cell(cell, name: str) -> None:
    _styled(cell, STATUS_STYLE.get(name, ("FFFFFF", "000000")))


def _done_means(entry: Entry, snap: Snapshot) -> str:
    """Each check with its mark, and the code tagged with that check, if any (file:line)."""
    states = snap.of(entry).criteria
    lines = []
    for c in entry.criteria:
        places = ", ".join(f"{f.path}:{line}" for f, line in snap.trace.code_for(entry.criterion_id(c.letter)))
        lines.append(f"{CHECK_MARK.get(states.get(c.letter, ''), '○')} {c.letter}. {c.text}"
                     + (f" — {places}" if places else ""))
    return "\n".join(lines)


def _signoffs(entry: Entry) -> str:
    return "\n".join(f"{s.on:%d %b %Y} · {s.what} · {s.person} · {s.note}" for s in entry.signoffs)


def _relative(snap: Snapshot, path: Path) -> str:
    return path.relative_to(snap.project.config.root).as_posix()


def _lines(text: str, width: int) -> int:
    """How many lines `text` wraps to in a column of `width` characters (viewers do not size rows reliably: P2)."""
    per_line = max(1, int(width * 1.1))
    return sum(max(1, math.ceil(len(part) / per_line)) for part in text.split("\n"))


def _table(ws: Worksheet, note: str, columns: list[Column], rows: list[list], *, styles: dict[int, dict] | None = None,
           links: dict[int, list[str]] | None = None, freeze: str = "B3") -> None:
    """A sheet with a do-not-edit note on row 1, a header on row 2 and wrapped, sized rows below.

    `styles` colours a column's cells by their value; `links` turns a column's cells into links (one per row).
    """
    ws.cell(row=1, column=1, value=note).font = NOTE_FONT
    ws.row_dimensions[1].height = LINE_HEIGHT + 3
    for col, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col)].width = column.width
        cell = ws.cell(row=2, column=col, value=column.title)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = LINE_HEIGHT * 2
    for index, values in enumerate(rows):
        r = index + 3
        lines = 1
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = BORDER
            lines = max(lines, _lines(str(value or ""), columns[c - 1].width))
        for col, style in (styles or {}).items():
            _styled(ws.cell(row=r, column=col), style.get(str(values[col - 1]), ("FFFFFF", "000000")))
        for col, targets in (links or {}).items():
            if targets[index]:
                cell = ws.cell(row=r, column=col)
                cell.hyperlink = targets[index]
                cell.font = LINK_FONT
        ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, LINE_HEIGHT * lines + 4)
    ws.freeze_panes = freeze
    if rows:
        ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{len(rows) + 2}"


def _section(ws: Worksheet, row: int, title: str) -> int:
    ws.cell(row=row, column=1, value=title).font = SECTION_FONT
    return row + 1


def _summary(ws: Worksheet, snap: Snapshot) -> None:
    project, trace = snap.project, snap.trace
    ws.column_dimensions["A"].width = 30
    for col in range(2, len(ORDER) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.cell(row=1, column=1, value=_do_not_edit(snap)).font = NOTE_FONT
    ws.cell(row=3, column=1, value=project.config.product).font = TITLE_FONT
    ws.row_dimensions[3].height = 24

    done, active = snap.proven_share()
    row = _section(ws, 5, "Features proven or live")
    ws.cell(row=row - 1, column=2, value=f"{done} of {active}" if active else "none agreed yet")
    row += 1

    for heading, entries, by_area in (("Features by area", project.features, True),
                                      ("Promises", project.promises, False)):
        row = _section(ws, row, heading)
        shown = [name for name in ORDER if snap.counts(entries).get(name)]
        for col, title in enumerate(["Area" if by_area else "", *shown, "Total"], start=1):
            cell = ws.cell(row=row, column=col, value=title)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, horizontal="center" if col > 1 else "left")
        row += 1
        groups: dict[str, list[Entry]] = {}
        if by_area:
            order = [*project.config.areas, *sorted({e.area for e in entries} - set(project.config.areas))]
            for area in order:
                members = [e for e in entries if e.area == area]
                if members:
                    groups[area or "(no area)"] = members
        else:
            groups["All promises"] = entries
        for label, members in groups.items():
            counts = snap.counts(members)
            ws.cell(row=row, column=1, value=label)
            for col, name in enumerate(shown, start=2):
                if counts[name]:
                    cell = ws.cell(row=row, column=col, value=counts[name])
                    _status_cell(cell, name)
                    cell.alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=len(shown) + 2, value=len(members)).alignment = Alignment(horizontal="center")
            row += 1
        if not entries:
            ws.cell(row=row, column=1, value="None yet").font = NOTE_FONT
            row += 1
        row += 1

    row = _section(ws, row, "Code")
    kinds = {kind: sum(1 for f in trace.files if f.kind == kind) for kind in CODE_STYLE}
    total = len(trace.files)
    for kind, count in kinds.items():
        ws.cell(row=row, column=1, value=kind)
        cell = ws.cell(row=row, column=2, value=count)
        cell.alignment = Alignment(horizontal="center")
        if kind != "Feature code" and count:
            _styled(cell, CODE_STYLE[kind])
        if total:
            ws.cell(row=row, column=3, value=f"{round(100 * count / total)}% of {total} files")
        row += 1
    row += 1

    row = _section(ws, row, "Test results")
    for results in trace.results.files:
        repo_name = next((r.name for r in trace.repos if r.repo == results.repo), results.repo.path.name)
        ws.cell(row=row, column=1, value=f"{repo_name}: {results.path.name}")
        ws.cell(row=row, column=2, value=f"{results.written:%d %b %Y %H:%M}")
        row += 1
    if not trace.results.files:
        ws.cell(row=row, column=1, value="None found: statuses do not include any tests yet. Set `results` in "
                                          "charter.toml to where your tests write JUnit XML.").font = NOTE_FONT
        row += 1
    row += 1

    row = _section(ws, row, "Changed in the last 7 days")
    by_id = {e.id: e for e in project.entries}
    for entry_id in snap.recent:
        entry = by_id[entry_id]
        ws.cell(row=row, column=1, value=f"{entry.id} · {entry.title}")
        _status_cell(ws.cell(row=row, column=2, value=snap.of(entry).name), snap.of(entry).name)
        row += 1
    if not snap.recent:
        ws.cell(row=row, column=1, value="Nothing").font = NOTE_FONT
        row += 1
    row += 1

    row = _section(ws, row, "Open task cards on this branch")
    open_cards = [c for c in project.cards if c.state == "open"]
    for card in open_cards:
        ws.cell(row=row, column=1, value=f"{card.id} · {card.title}")
        ws.cell(row=row, column=2, value=", ".join(card.serves))
        ws.cell(row=row, column=4, value=f"Next: {card.next_step}" if card.next_step else "")
        row += 1
    if not open_cards:
        ws.cell(row=row, column=1, value="None").font = NOTE_FONT
        row += 1
    row += 1

    row = _section(ws, row, "Work in progress on other branches")
    for work in snap.wip.branches:
        ws.cell(row=row, column=1, value=f"{work.branch.name} ({work.branch.author}, {work.when:%d %b})")
        ws.cell(row=row, column=2, value=", ".join(f"{c.id} {c.title}" for c in work.cards))
        row += 1
    if not snap.wip.branches:
        ws.cell(row=row, column=1, value="None found" + (f" (compared with {snap.wip.base})" if snap.wip.base
                                                          else "")).font = NOTE_FONT


def _tag_line(snap: Snapshot, tag) -> str:
    """One tag as a person reads it: what it serves, in the charter's words, and where."""
    entries = {e.id: e for e in snap.project.entries}
    names = []
    for ref in tag.refs:
        entry_id, _, letter = ref.partition(".")
        entry = entries.get(entry_id)
        if entry is None:
            names.append(f"{ref} · not in the charter")
        elif letter:
            text = next((c.text for c in entry.criteria if c.letter == letter), "no such check")
            names.append(f"{ref} · {text}")
        else:
            names.append(f"{ref} · {entry.title}")
    return "; ".join(names) + f" (line {tag.line})"


def _code_map(ws: Worksheet, snap: Snapshot) -> None:
    files = sorted(snap.trace.files, key=lambda f: (f.repo.name, f.path))
    rows = [[f.repo.name, f.path, f.kind, "\n".join(_tag_line(snap, tag) for tag in f.tags if tag.refs)]
            for f in files]
    links = [f.link for f in files]
    commits = sorted({f"{r.name} at {r.commit[:7]}" for r in snap.trace.repos if r.commit})
    unpushed = sorted(r.name for r in snap.trace.repos if r.commit and not r.pushed)
    note = _do_not_edit(snap)
    if commits:
        note += f" File links open {', '.join(commits)}."
    if unpushed:
        note += f" Not pushed yet, so no links: {', '.join(unpushed)}."
    _table(ws, note, [Column("Repository", 18), Column("File", 52), Column("What it is", 18),
                      Column("Serves", 46)], rows,
           styles={3: CODE_STYLE}, links={2: links}, freeze="C3")


def _work_in_progress(ws: Worksheet, snap: Snapshot) -> None:
    wip = snap.wip
    rows = [
        [work.branch.name, "yes" if work.branch.pushed else "not yet", work.branch.author,
         f"{work.when:%d %b %Y %H:%M}", work.ahead, card.id, card.title, ", ".join(card.serves), card.next_step]
        for work in wip.branches for card in work.cards
    ]
    note = _do_not_edit(snap)
    if wip.base:
        note += f" Open task cards on branches with work not yet merged into {wip.base}"
        note += f", as of the last fetch on {wip.fetched:%d %b %H:%M}." if wip.fetched else "."
    _table(ws, note, [Column("Branch", 30), Column("Pushed", 9), Column("Last commit by", 18), Column("When", 17),
                      Column("Commits not merged", 11), Column("Card", 10), Column("Title", 30), Column("Serves", 22),
                      Column("Next step", 50)], rows, freeze="B3")


def write_xlsx(snap: Snapshot, path: Path) -> None:
    project = snap.project
    note = _do_not_edit(snap)
    titles = {d.id: d.title for d in project.decisions}
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _summary(summary, snap)

    features = sorted(project.features, key=lambda e: e.id)
    sheet = wb.create_sheet("Features")
    _table(sheet, _yours(snap), [
        Column("ID", 8), Column("Title", 22), Column(DECISION_HEADER, 16), Column(NOTE_HEADER, 28),
        Column("Area", 16), Column("Priority", 9), Column("Release", 12), Column("Status", 12), Column("Checks", 14),
        Column("Built in", 34), Column("Proven by", 30), Column("As a", 16), Column("I want to", 28),
        Column("With", 32), Column("Must also", 24), Column("So that", 28), Column("Done means", 44),
        Column("Sign-offs", 50), Column("Decisions", 24), Column("File", 30),
    ], [
        [f.id, f.title, current(f), "", f.area, f.priority, f.release, snap.of(f).name, snap.of(f).checks,
         snap.built_in(f), snap.proven_by(f), f.story.get("As a", ""), f.story.get("I want to", ""),
         f.story.get("With", ""), f.story.get("Must also", ""), f.story.get("So that", ""), _done_means(f, snap),
         _signoffs(f),
         "\n".join(f"{d} · {titles.get(d, 'not found')}" for d in f.decisions), _relative(snap, f.path)]
        for f in features
    ], styles={8: STATUS_STYLE}, links={20: [_relative(snap, f.path) for f in features]})
    _editable(sheet, 3, 4, choices(FEATURE), len(features))

    promises = sorted(project.promises, key=lambda e: e.id)
    sheet = wb.create_sheet("Promises")
    _table(sheet, _yours(snap), [
        Column("ID", 8), Column("Title", 24), Column(DECISION_HEADER, 14), Column(NOTE_HEADER, 28),
        Column("Status", 12), Column("Checks", 14), Column("Built in", 34), Column("Proven by", 30),
        Column("We promise", 40), Column("Applies to", 28), Column("Done means", 44), Column("Sign-offs", 50),
        Column("File", 30),
    ], [
        [p.id, p.title, current(p), "", snap.of(p).name, snap.of(p).checks, snap.built_in(p), snap.proven_by(p),
         p.story.get("We promise", ""), p.story.get("Applies to", ""), _done_means(p, snap), _signoffs(p),
         _relative(snap, p.path)]
        for p in promises
    ], styles={5: STATUS_STYLE}, links={13: [_relative(snap, p.path) for p in promises]})
    _editable(sheet, 3, 4, choices(PROMISE), len(promises))

    _code_map(wb.create_sheet("Code map"), snap)
    _work_in_progress(wb.create_sheet("Work in progress"), snap)

    decisions = sorted(project.decisions, key=lambda d: d.id)
    _table(wb.create_sheet("Decisions"), note, [
        Column("ID", 9), Column("Date", 12), Column("Title", 26), Column("Decision", 40), Column("Instead of", 30),
        Column("Because", 40), Column("Decided by", 16), Column("Features", 16), Column("File", 30),
    ], [
        [d.id, f"{d.on:%d %b %Y}" if d.on else "", d.title, d.decision, d.instead_of, d.because, d.decided_by,
         ", ".join(d.features), _relative(snap, d.path)]
        for d in decisions
    ], links={9: [_relative(snap, d.path) for d in decisions]})

    _table(wb.create_sheet("Ideas"), note, [Column("Date", 12), Column("Idea", 70), Column("From", 30)],
           [[f"{i.on:%d %b %Y}" if i.on else "", i.text, i.source] for i in project.ideas], freeze="A3")

    wb.save(path)
